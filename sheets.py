"""Google Sheets sync through a Google Apps Script web app (no Google client libs / credentials in this repo).

push: POST {secret, sheets:{tab:[[header..],[row..]]}}  -> script rewrites each tab
pull: GET  ?secret=...                                  -> script returns the same structure
Also offers an .xlsx export (importable into Google Sheets) as an offline fallback.
"""
import os

import requests
from openpyxl import Workbook

from db import PROFILE_FIELDS, TABLES

BASE = os.path.dirname(os.path.abspath(__file__))


def configured():
    return bool(os.environ.get("SHEETS_WEBAPP_URL"))


def to_tabs(pf):
    tabs = {"profile": [PROFILE_FIELDS, [[pf["profile"].get(k, "") for k in PROFILE_FIELDS]] if pf["profile"] else []]}
    for t, fields in TABLES.items():
        tabs[t] = [["id"] + fields, [[r["id"]] + [r.get(f, "") for f in fields] for r in pf.get(t, [])]]
    return {name: [hdr] + rows for name, (hdr, rows) in tabs.items()}


def from_tabs(tabs):
    pf = {"profile": {}}
    prof = tabs.get("profile") or []
    if len(prof) > 1:
        pf["profile"] = dict(zip(prof[0], prof[1]))
    for t, fields in TABLES.items():
        grid = tabs.get(t) or []
        pf[t] = []
        if len(grid) > 1:
            hdr = grid[0]
            for row in grid[1:]:
                rec = dict(zip(hdr, row))
                if str(rec.get(fields[0], "")).strip() or str(rec.get("company", "")).strip():
                    pf[t].append({f: ("" if rec.get(f) is None else str(rec.get(f))) for f in fields})
    return pf


def _url():
    return os.environ["SHEETS_WEBAPP_URL"]


def push(pf):
    if not configured():
        raise RuntimeError("SHEETS_WEBAPP_URL not set; use the .xlsx export or see README to deploy Code.gs")
    r = requests.post(_url(), json={"secret": os.environ.get("SHEETS_SECRET", ""), "sheets": to_tabs(pf)},
                      timeout=60, allow_redirects=True)
    r.raise_for_status()
    res = r.json()
    if not res.get("ok"):
        raise RuntimeError(res.get("error", "sheet push failed"))
    return {t: max(len(g) - 1, 0) for t, g in to_tabs(pf).items()}


def pull():
    if not configured():
        raise RuntimeError("SHEETS_WEBAPP_URL not set")
    r = requests.get(_url(), params={"secret": os.environ.get("SHEETS_SECRET", "")}, timeout=60)
    r.raise_for_status()
    res = r.json()
    if not res.get("ok"):
        raise RuntimeError(res.get("error", "sheet pull failed"))
    return from_tabs(res["sheets"])


def export_xlsx(pf, path=None):
    path = path or os.path.join(BASE, "exports", "portfolio_sheet.xlsx")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name, grid in to_tabs(pf).items():
        ws = wb.create_sheet(name)
        for row in grid:
            ws.append(row)
        for cell in ws[1]:
            cell.font = cell.font.copy(bold=True)
    wb.save(path)
    return path
